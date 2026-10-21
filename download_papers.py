import csv
import re
import sys
from pathlib import Path

import gdown
import openpyxl

xlsx = sys.argv[1]
out = Path(sys.argv[2] if len(sys.argv) > 2 else "papers")
out.mkdir(parents=True, exist_ok=True)

ws = openpyxl.load_workbook(xlsx)["1_Screening"]
header = [c.value for c in ws[1]]
col = {name: i for i, name in enumerate(header) if name}

rows = []
for r in range(2, ws.max_row + 1):
    cell = ws.cell(r, 5)
    if not (cell.hyperlink and "drive.google.com" in cell.hyperlink.target):
        continue
    m = re.search(r"/d/([\w-]+)", cell.hyperlink.target)
    get = lambda name: ws.cell(r, col[name] + 1).value if name in col else None
    study_id = str(get("Study\\_ID") or get("Study_ID") or f"row{r}").strip()
    safe = re.sub(r"[^\w.-]+", "_", study_id)
    dest = out / f"{safe}.pdf"
    status = "exists"
    if not dest.exists():
        try:
            ok = gdown.download(id=m.group(1), output=str(dest), quiet=True)
            status = "ok" if ok else "failed"
        except Exception as e:  # permission denied, quota, etc.
            status = f"failed: {type(e).__name__}"
            dest.unlink(missing_ok=True)
    print(f"{status:8} {dest.name}")
    rows.append({
        "row": r,
        "study_id": study_id,
        "file": dest.name,
        "status": status,
        "title": get("Title"),
        "year": get("Year"),
        "final_decision": get("Final Decision (I/E/M)"),
        "drive_url": cell.hyperlink.target,
    })

with open(out / "manifest.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=rows[0].keys())
    w.writeheader()
    w.writerows(rows)

failed = [x for x in rows if x["status"].startswith("failed")]
print(f"\n{len(rows) - len(failed)}/{len(rows)} downloaded; {len(failed)} failed")
